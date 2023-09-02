from copy import deepcopy

from docx.oxml import CT_Tbl, CT_R
from docx.shared import Parented, Length
from docx.table import Table, _Row, _Cell

from md2gost.util import create_element


__all__ = [
    "create_table",
    "create_table_row",
    "create_table_cell",
]


def create_table(parent: Parented, rows: int, cols, width: Length, style="Table Grid"):
    table = Table(CT_Tbl.new_tbl(rows, cols, width), parent)
    table.style = style

    # google docs fix
    table._tbl.tblPr.append(
        deepcopy(parent.part.styles[style]._element.xpath("w:tblPr/w:tblBorders")[0]))
    for i in range(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell._element.remove(cell.paragraphs[0]._element)
            table.cell(i, j)._element.tcPr.append(create_element("w:shd", {
                "w:fill": "auto", "w:val": "clear"
            }))

    return table


def create_table_row(parent: Table):
    return _Row(create_element("w:tr"), parent)


def create_table_cell(parent: _Row, width: Length):
    cell = _Cell(create_element("w:tc"), parent)
    cell.width = width
    return cell


def create_field(text: str, instr_text: str) -> CT_R:
    r = create_element("w:r")

    r.append(create_element("w:fldChar", {
        "w:fldCharType": "begin"
    }))
    r.append(create_element("w:instrText", {
        "xml:space": "preserve"
    }, instr_text))
    r.append(create_element("w:fldChar", {
        "w:fldCharType": "separate"
    }))
    if text is not None:
        r.append(create_element("w:t", text))
    r.append(create_element("w:fldChar", {
        "w:fldCharType": "end"
    }))

    return r
